"""Build the bundled Cigar Market Map catalog from the validated research workbook."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "outputs" / "market_intel" / "Mapa_Mundial_Puros_Mercados_Distribuidores_Tiendas.xlsx"
OUTPUT = ROOT / "web" / "data" / "locations.json"
PENDING = {"", "pendiente", "n/a", "na", "none", "null"}


def text(value):
    value = str(value or "").strip()
    return "" if value.casefold() in PENDING else value


def key(*values):
    joined = "|".join(text(value).casefold() for value in values)
    return re.sub(r"[^a-z0-9]+", "-", joined).strip("-")[:180]


def rows(sheet):
    headers = [text(cell.value) for cell in sheet[6]]
    for values in sheet.iter_rows(min_row=7, values_only=True):
        record = dict(zip(headers, values))
        if any(text(value) for value in values):
            yield record


def store_record(row, source_label):
    name = text(row.get("Nombre"))
    country = text(row.get("Pais"))
    city = text(row.get("Ciudad"))
    address = text(row.get("Direccion"))
    notes = " | ".join(filter(None, [
        f"Autorizacion: {text(row.get('Autorizacion'))}" if text(row.get("Autorizacion")) else "",
        f"Actividad: {text(row.get('Actividad'))}" if text(row.get("Actividad")) else "",
        text(row.get("Notas mapa")),
    ]))
    return {
        "id": f"store-{key(name, address, city, country)}",
        "type": text(row.get("Tipo")) or "Tienda",
        "relevance": text(row.get("Relevancia")) or "Sin clasificar",
        "dataStatus": text(row.get("Estado dato")) or "Pendiente",
        "name": name or "Registro sin nombre",
        "legalName": "",
        "address": address,
        "city": city,
        "province": "",
        "country": country,
        "postalCode": "",
        "lat": None,
        "lng": None,
        "phone": text(row.get("Telefono")),
        "email": text(row.get("Email")),
        "web": " | ".join(filter(None, [text(row.get("Web")), text(row.get("Instagram")), text(row.get("Facebook"))])),
        "contact": text(row.get("Distribuidor")),
        "status": "Sin contactar",
        "nextAction": "Validar contacto y geocodificar" if not address else "Validar contacto",
        "notes": notes,
        "source": text(row.get("Fuente")) or source_label,
        "sourceDate": "2026-09-15",
    }


def distributor_record(row):
    name = text(row.get("Nombre"))
    country = text(row.get("Pais"))
    city = text(row.get("Ciudad"))
    address = text(row.get("Direccion"))
    notes = " | ".join(filter(None, [
        f"Autorizacion: {text(row.get('Autorizacion'))}" if text(row.get("Autorizacion")) else "",
        f"Alcance: {text(row.get('Marcas/Alcance'))}" if text(row.get("Marcas/Alcance")) else "",
        text(row.get("Notas")),
    ]))
    return {
        "id": f"distributor-{key(name, address, city, country)}",
        "type": text(row.get("Tipo actor")) or "Distribuidor",
        "relevance": text(row.get("Relevancia")) or "Sin clasificar",
        "dataStatus": text(row.get("Estado dato")) or "Pendiente",
        "name": name or "Distribuidor sin nombre",
        "legalName": name,
        "address": address,
        "city": city,
        "province": "",
        "country": country,
        "postalCode": "",
        "lat": None,
        "lng": None,
        "phone": text(row.get("Telefono")),
        "email": text(row.get("Email")),
        "web": " | ".join(filter(None, [text(row.get("Web")), text(row.get("Redes"))])),
        "contact": "",
        "status": "Sin contactar",
        "nextAction": "Validar contacto comercial",
        "notes": notes,
        "source": text(row.get("Fuente")) or "Investigacion distribuidores",
        "sourceDate": "2026-09-15",
    }


def contact_record(row):
    company = text(row.get("Empresa"))
    contact = text(row.get("Nombre contacto"))
    country = text(row.get("Pais/Region"))
    city = text(row.get("Ciudad"))
    return {
        "id": f"contact-{key(company, contact, city, country)}",
        "type": "Contacto personal",
        "relevance": "Alta" if "prioridad" in text(row.get("Notas")).casefold() else "Media",
        "dataStatus": text(row.get("Estado dato")) or "Pendiente",
        "name": contact or company or "Contacto sin nombre",
        "legalName": company,
        "address": "",
        "city": city,
        "province": "",
        "country": country,
        "postalCode": "",
        "lat": None,
        "lng": None,
        "phone": text(row.get("Telefono")),
        "email": text(row.get("Email")),
        "web": " | ".join(filter(None, [text(row.get("Web")), text(row.get("LinkedIn/Redes"))])),
        "contact": f"{contact} - {text(row.get('Cargo/rol'))}".strip(" -"),
        "status": "Sin contactar",
        "nextAction": "Contacto B2B a validar",
        "notes": text(row.get("Notas")),
        "source": text(row.get("Fuente")) or "Investigacion contactos",
        "sourceDate": "2026-09-15",
    }


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    output = []
    seen = set()

    for sheet_name, label in [("Tiendas y Lounges", "Tiendas y Lounges"), ("Tiendas AUTO 428", "Tiendas AUTO 428")]:
        for row in rows(workbook[sheet_name]):
            record = store_record(row, label)
            fingerprint = key(record["type"], record["name"], record["address"], record["city"], record["country"])
            if fingerprint not in seen:
                output.append(record)
                seen.add(fingerprint)

    for row in rows(workbook["Distribuidores"]):
        record = distributor_record(row)
        output.append(record)

    for row in rows(workbook["Contactos Clave"]):
        record = contact_record(row)
        output.append(record)

    counts = {}
    for record in output:
        counts[record["type"]] = counts.get(record["type"], 0) + 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({"generatedAt": "2026-09-15", "source": SOURCE.name, "counts": counts, "records": output}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"records": len(output), "counts": counts}, ensure_ascii=False))


if __name__ == "__main__":
    main()
